"""claw xlsx chart — add a chart to a sheet."""

from __future__ import annotations

from pathlib import Path

import click

from claw.common import EXIT_INPUT, common_output_options, die, emit_json, safe_write


CHART_CLS = {
    "bar":     ("openpyxl.chart", "BarChart"),
    "col":     ("openpyxl.chart", "BarChart"),
    "line":    ("openpyxl.chart", "LineChart"),
    "pie":     ("openpyxl.chart", "PieChart"),
    "scatter": ("openpyxl.chart", "ScatterChart"),
    "area":    ("openpyxl.chart", "AreaChart"),
}


@click.command(name="chart")
@click.argument("src", type=click.Path(exists=True, path_type=Path))
@click.option("--sheet", required=True)
@click.option("--type", "chart_type", required=True, type=click.Choice(list(CHART_CLS.keys())))
@click.option("--data", "data_range", required=True, help="A1 range for series data.")
@click.option("--categories", "cat_range", default=None, help="A1 range for categories.")
@click.option("--title", default=None)
@click.option("--x-title", "x_title", default=None)
@click.option("--y-title", "y_title", default=None)
@click.option("--at", "at_cell", default="F2")
@click.option("--style", "style_n", default=None, type=int)
@common_output_options
def chart(src: Path, sheet: str, chart_type: str, data_range: str,
          cat_range: str | None, title: str | None,
          x_title: str | None, y_title: str | None, at_cell: str, style_n: int | None,
          force: bool, backup: bool, as_json: bool, dry_run: bool,
          quiet: bool, verbose: bool, mkdir: bool) -> None:
    """Add a bar / col / line / pie / scatter / area chart."""
    try:
        import importlib
        from openpyxl import load_workbook
        from openpyxl.chart import Reference
    except ImportError:
        die("openpyxl not installed", code=EXIT_INPUT,
            hint="uv tool install 'claw[xlsx]'", as_json=as_json)

    if dry_run:
        click.echo(f"would add {chart_type} chart at {at_cell} on {sheet}")
        return

    wb = load_workbook(src)
    ws = wb[sheet]

    mod_name, cls_name = CHART_CLS[chart_type]
    cls = getattr(importlib.import_module(mod_name), cls_name)
    ch = cls()
    if chart_type == "col":
        ch.type = "col"
    elif chart_type == "bar":
        ch.type = "bar"
    if title:
        ch.title = title
    if x_title and hasattr(ch, "x_axis"):
        ch.x_axis.title = x_title
    if y_title and hasattr(ch, "y_axis"):
        ch.y_axis.title = y_title
    if style_n is not None:
        ch.style = style_n

    data_ref = Reference(ws, range_string=f"{sheet}!{data_range}")
    ch.add_data(data_ref, titles_from_data=True)
    if cat_range:
        cat_ref = Reference(ws, range_string=f"{sheet}!{cat_range}")
        ch.set_categories(cat_ref)

    ws.add_chart(ch, at_cell)

    def _save(f):
        wb.save(f)

    safe_write(src, _save, force=True, backup=backup, mkdir=mkdir)

    if as_json:
        emit_json({"path": str(src), "sheet": sheet, "chart": chart_type, "at": at_cell})
    elif not quiet:
        click.echo(f"added {chart_type} chart at {at_cell} on {sheet}")
